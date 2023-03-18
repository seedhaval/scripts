import openpyxl

from lib.codehelper import get_safe_output_xls_path, data_path, \
    get_column_config_for_subject
from lib.db import get_exam_map_for_all_subjects
from lib.excelhelper import save_close_and_start, add_table
import json


def get_column_ids():
    filepath = get_safe_output_xls_path("column_id", False)
    wb = openpyxl.Workbook()
    sht = wb.active
    with open(data_path + "\\export_columns.json", encoding='utf8') as f:
        subs = json.load(f).keys()
    exam_map = get_exam_map_for_all_subjects()
    out = [["Subject", "Exam", "Column ID"]]
    for sub in subs:
        colinfo = get_column_config_for_subject(sub)
        for col in colinfo:
            colnm = col['nm'] if 'nm' in col else exam_map[int(col['id'])][0]
            out.append([sub, colnm, str(col['id'])])
    add_table(sht, 1, 1, out)
    save_close_and_start(wb, filepath)


#def csv_to_json_col_export():
#    d = {}
#    with open(data_path + "\\export_columns", encoding='utf8') as f:
#        sub_ar = [x.strip().split(",")[0] for x in f.readlines() if x.strip()]
#    for sub in sub_ar:
#        curd = get_column_config_for_subject(sub)
#        for row in curd:
#            del row["color"]
#        d[sub] = curd
#
#    print(json.dumps(d, indent=2, ensure_ascii=False))
